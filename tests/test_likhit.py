"""
test_likhit.py - poore pipeline ki jaanch, bina asli API key ke.

Kaise: cloud aur local servers ki jagah chhote nakli (mock) HTTP servers
chalaate hain jo asli servers jaisa jawab dete hain. Isse check hota hai ki
hamari request sahi format mein jaa rahi hai aur jawab sahi padha ja raha hai.
Tesseract wala test asli Tesseract chalata hai (agar install ho).

Chalane ka tareeka:  python -m unittest discover -s tests -v
"""

import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from likhit import engines  # noqa: E402
from likhit.docx_writer import Page, build_document  # noqa: E402
from likhit.settings import Settings  # noqa: E402


def sample_image(text="Hello from Likhit\nThis is a test page"):
    """Test ke liye ek safed kaagaz par likha hua text."""
    img = Image.new("RGB", (1400, 500), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 56)
    except OSError:
        font = ImageFont.load_default()
    y = 60
    for line in text.splitlines():
        draw.text((60, y), line, fill="black", font=font)
        y += 110
    return img


class FakeServer:
    """Ek chhota HTTP server jo har path ke liye diya hua JSON lautata hai
    aur aayi hui requests yaad rakhta hai."""

    def __init__(self, routes):
        self.requests = []
        outer = self

        class Handler(BaseHTTPRequestHandler):
            def _reply(self):
                length = int(self.headers.get("Content-Length") or 0)
                body = json.loads(self.rfile.read(length) or b"null") if length else None
                outer.requests.append((self.path, dict(self.headers), body))
                handler = routes.get(self.path.split("?")[0])
                if handler is None:
                    self.send_response(404)
                    self.end_headers()
                    return
                data = json.dumps(handler(body)).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

            do_GET = do_POST = _reply

            def log_message(self, *a):
                pass

        self.httpd = HTTPServer(("127.0.0.1", 0), Handler)
        self.url = f"http://127.0.0.1:{self.httpd.server_address[1]}"
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()

    def close(self):
        self.httpd.shutdown()


class CleanOutputTests(unittest.TestCase):
    def test_strips_fences_and_preamble(self):
        raw = "```text\nHere is the transcription:\nDear diary\n```"
        self.assertEqual(engines.clean_output(raw), "Dear diary")

    def test_keeps_normal_text(self):
        self.assertEqual(engines.clean_output("  Aaj ka din  "), "Aaj ka din")


class DocxTests(unittest.TestCase):
    def test_formatting_and_watermark(self):
        from docx import Document
        text = ("# My Notes\n\nFirst line of a\nparagraph that wraps.\n\n"
                "- apples\n- mangoes\n\nThis word is **important** and this one unclear[?]")
        with tempfile.TemporaryDirectory() as tmp:
            out = build_document([Page("p1", sample_image(), text)], Path(tmp) / "t.docx",
                                 include_scans=True)
            doc = Document(out)
            styles = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
            self.assertIn(("Heading 1", "My Notes"), styles)
            self.assertIn(("Normal", "First line of a paragraph that wraps."), styles)
            self.assertIn(("List Bullet", "apples"), styles)
            bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
            self.assertIn("important", bold)
            hl = [r.text for p in doc.paragraphs for r in p.runs if r.font.highlight_color]
            self.assertIn("[?]", hl)
            # Header mein logo image honi chahiye (watermark)
            header_xml = doc.sections[0].header._element.xml
            self.assertIn("<pic:pic", header_xml)
            self.assertIn("Kanak Prabhakar", doc.sections[0].header.paragraphs[0].text)
            # include_scans=True: body mein bhi ek image
            self.assertEqual(len(doc.inline_shapes), 1)


class EngineChoiceTests(unittest.TestCase):
    def found(self, **kw):
        base = {"cloud": False, "ollama": [], "lmstudio": [], "tesseract": ""}
        base.update(kw)
        return base

    def test_auto_prefers_cloud_when_key_exists(self):
        s = Settings(api_key="sk-test")
        e = engines.choose_engine(s, "auto", self.found(ollama=["llava"]))
        self.assertEqual(e.kind, "cloud")

    def test_auto_falls_back_to_local(self):
        e = engines.choose_engine(Settings(), "auto",
                                  self.found(ollama=["qwen2.5vl:3b"], tesseract="/x"))
        self.assertEqual(e.kind, "local")
        self.assertIn("qwen2.5vl", e.label)

    def test_auto_falls_back_to_offline(self):
        if not engines.find_tesseract():
            self.skipTest("tesseract not installed")
        e = engines.choose_engine(Settings(), "auto",
                                  self.found(tesseract=engines.find_tesseract()))
        self.assertEqual(e.kind, "offline")

    def test_nothing_available_explains_all_three(self):
        with self.assertRaises(engines.EngineError) as ctx:
            engines.choose_engine(Settings(), "auto", self.found())
        msg = str(ctx.exception)
        self.assertIn("API key", msg)
        self.assertIn("ollama", msg.lower())
        self.assertIn("Tesseract", msg)

    def test_forced_mode_does_not_fall_back(self):
        with self.assertRaises(engines.EngineError):
            engines.choose_engine(Settings(), "cloud", self.found(ollama=["llava"]))


class LocalOllamaTests(unittest.TestCase):
    def setUp(self):
        self.server = FakeServer({
            "/api/tags": lambda b: {"models": [{"name": "llama3:8b"},
                                               {"name": "qwen2.5vl:3b"}]},
            "/api/show": lambda b: {"capabilities": ["completion", "vision"]
                                    if "vl" in b["model"] else ["completion"]},
            "/api/chat": lambda b: {"message": {"content": "```\nNamaste duniya\n```"}},
        })

    def tearDown(self):
        self.server.close()

    def test_detects_only_vision_models_and_reads(self):
        models = engines.detect_ollama(self.server.url)
        self.assertEqual(models, ["qwen2.5vl:3b"])
        engine = engines.OllamaEngine(self.server.url, models[0])
        self.assertEqual(engine.recognize(sample_image()), "Namaste duniya")
        path, _, body = self.server.requests[-1]
        self.assertEqual(path, "/api/chat")
        self.assertEqual(len(body["messages"][0]["images"]), 1)


class CloudTests(unittest.TestCase):
    def test_openai_compatible_request_shape(self):
        server = FakeServer({"/v1/chat/completions": lambda b: {
            "choices": [{"message": {"content": "Dear Sir,\n\nThank you."}}]}})
        try:
            s = Settings(cloud_provider="openai_compatible", api_key="k-123",
                         cloud_model="some-vision-model", custom_base_url=server.url + "/v1")
            engine = engines.choose_engine(s, "cloud", {"cloud": True, "ollama": [],
                                                        "lmstudio": [], "tesseract": ""})
            self.assertEqual(engine.recognize(sample_image()), "Dear Sir,\n\nThank you.")
            _, headers, body = server.requests[-1]
            self.assertEqual(headers.get("Authorization"), "Bearer k-123")
            parts = body["messages"][0]["content"]
            self.assertTrue(parts[1]["image_url"]["url"].startswith("data:image/jpeg;base64,"))
        finally:
            server.close()

    def test_bad_key_gives_clear_message(self):
        class Deny(BaseHTTPRequestHandler):
            def do_POST(self):
                self.send_response(401)
                self.end_headers()

            def log_message(self, *a):
                pass
        httpd = HTTPServer(("127.0.0.1", 0), Deny)
        threading.Thread(target=httpd.serve_forever, daemon=True).start()
        try:
            url = f"http://127.0.0.1:{httpd.server_address[1]}"
            with self.assertRaises(engines.EngineError) as ctx:
                engines.openai_style_chat(url, "m", "AAAA")
            self.assertIn("API key", str(ctx.exception))
        finally:
            httpd.shutdown()


class ProviderFormatTests(unittest.TestCase):
    """Anthropic aur Gemini ke asli URL par nahi jaate; _http_json ko
    thodi der ke liye badal kar dekhte hain ki request sahi bani."""

    def capture(self, reply):
        calls = []
        original = engines._http_json
        engines._http_json = lambda url, payload=None, headers=None, timeout=0: (
            calls.append((url, payload, headers)) or reply)
        self.addCleanup(setattr, engines, "_http_json", original)
        return calls

    def test_anthropic(self):
        calls = self.capture({"content": [{"type": "text", "text": "Likha hua"}]})
        out = engines.CloudEngine("anthropic", "key-a").recognize(sample_image())
        url, payload, headers = calls[0]
        self.assertEqual(out, "Likha hua")
        self.assertTrue(url.endswith("/v1/messages"))
        self.assertEqual(headers["x-api-key"], "key-a")
        self.assertEqual(payload["messages"][0]["content"][0]["type"], "image")

    def test_gemini(self):
        calls = self.capture({"candidates": [{"content": {"parts": [{"text": "Notes"}]}}]})
        out = engines.CloudEngine("gemini", "key-g").recognize(sample_image())
        url, payload, headers = calls[0]
        self.assertEqual(out, "Notes")
        self.assertIn(":generateContent", url)
        self.assertEqual(headers["x-goog-api-key"], "key-g")
        self.assertIn("inline_data", payload["contents"][0]["parts"][1])


class PdfTests(unittest.TestCase):
    def test_pdf_pages_load(self):
        from likhit import imaging
        if not imaging.PDF_SUPPORTED:
            self.skipTest("pymupdf not installed")
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "two.pdf"
            sample_image().save(pdf, save_all=True, append_images=[sample_image("Page two")])
            pages = imaging.load_pages(str(pdf))
            self.assertEqual(len(pages), 2)
            self.assertEqual(pages[1][0], "two - page 2")


@unittest.skipUnless(engines.find_tesseract(), "tesseract not installed")
class OfflineTests(unittest.TestCase):
    def test_tesseract_reads_text(self):
        engine = engines.TesseractEngine(engines.find_tesseract(), "eng")
        text = engine.recognize(sample_image()).lower()
        self.assertIn("hello", text)
        self.assertIn("test page", text)

    def test_cli_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            img_path = Path(tmp) / "page.png"
            sample_image().save(img_path)
            out = Path(tmp) / "out.docx"
            env = dict(os.environ, XDG_CONFIG_HOME=tmp, APPDATA=tmp)
            r = subprocess.run([sys.executable, str(ROOT / "main.py"), str(img_path),
                                "-o", str(out), "--mode", "offline"],
                               capture_output=True, text=True, env=env, timeout=120)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            self.assertTrue(out.exists())


if __name__ == "__main__":
    unittest.main()
