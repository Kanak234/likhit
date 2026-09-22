"""
docx_writer.py - padha hua text asli Word (.docx) file mein badalna.

Why: sirf text copy karke Word mein paste karna kaafi nahi. Headings asli
Heading style mein hon, bullets asli list hon, taaki user Word mein Table of
Contents bana sake aur formatting ek jhatke mein badal sake. Jo shabd AI nahi
padh paaya ([?]) wo peele highlight mein dikhta hai, taaki check karna aasaan
ho.

Flow: ui.py / main.py -> build_document(pages, path) -> .docx disk par.
Har page ke upar header mein KP logo (Kanak Prabhakar ka watermark) hota hai.
"""

import io
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import imaging

ASSETS = Path(__file__).resolve().parent.parent / "assets"
# Word ka page safed hota hai, isliye dark orbit line wala logo
LOGO_FOR_DOCX = ASSETS / "logo-light-bg-128.png"


@dataclass(eq=False)  # eq=False: pages ko identity se pehchano, pixel compare nahi
class Page:
    """Ek photo/PDF page aur uska padha hua text. ui.py isi ki list rakhta hai."""
    name: str
    image: object          # PIL Image
    text: str = ""
    engine_label: str = ""  # kis engine ne padha, sirf status ke liye
    # v1.1: page padhne mein fail hua toh wajah yahan. Word file mein us page
    # ki jagah ek note aata hai, taaki page gayab na ho aur kram na bigde.
    error: str = ""


NO_TEXT = "[no text found]"


# ---------------------------------------------------------------------------
# Grey, italic note: khaali ya fail hue page ki jagah.
# Why: pehle "[no text found]" jaisa marker text seedha Word mein chhap jaata
# tha, jo asli likhaai jaisa lagta tha. Note saaf dikhata hai ki ye Likhit
# ka sandesh hai, user ka likha hua nahi.
# ---------------------------------------------------------------------------
def _add_note(doc, text):
    run = doc.add_paragraph().add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)


# ---------------------------------------------------------------------------
# Ek line ke andar ki formatting: **bold** aur [?].
# Why: regex se line ko tukdon mein todte hain; har tukda Word ka ek "run"
# banta hai jiska apna style ho sakta hai.
# ---------------------------------------------------------------------------
_INLINE = re.compile(r"(\*\*.+?\*\*|\[\?\])")


def _add_inline(paragraph, text):
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece == "[?]":
            run = paragraph.add_run("[?]")
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            paragraph.add_run(piece[2:-2]).bold = True
        else:
            paragraph.add_run(piece)


# ---------------------------------------------------------------------------
# Page ka poora text Word paragraphs mein.
# Rule: khaali line = naya paragraph. "# " heading, "- " bullet, "1. " number.
# Lagaatar saadi lines ek hi paragraph mein jud jaati hain (handwriting mein
# line kaagaz khatam hone se tootti hai, matlab khatam hone se nahi).
# ---------------------------------------------------------------------------
def _write_text(doc, text):
    buffer = []

    def flush():
        if buffer:
            _add_inline(doc.add_paragraph(), " ".join(buffer))
            buffer.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            flush()
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            # v1.1 fix: pehle sirf "# " aur "## " pehchaane jaate the; "### "
            # hashes ke saath saada text ban jaata tha. Ab 1 se 6 tak sab
            # Word ke Heading 1-6 ban jaate hain.
            flush()
            level = len(heading.group(1))
            _add_inline(doc.add_heading(level=level), heading.group(2).strip())
        elif re.match(r"^[-*•]\s+", line):
            flush()
            _add_inline(doc.add_paragraph(style="List Bullet"),
                        re.sub(r"^[-*•]\s+", "", line))
        elif re.match(r"^\d+[.)]\s+", line):
            flush()
            # Number ko text mein hi rakhte hain; Word ki auto-numbering har
            # page par 1 se dobara shuru nahi hoti, jo handwriting se alag hota
            _add_inline(doc.add_paragraph(), line)
        else:
            buffer.append(line)
    flush()


# ---------------------------------------------------------------------------
# Header (logo watermark) aur footer (page number).
# Why: section header har page par apne aap aata hai; ek baar lagao, poore
# document mein dikhega.
# ---------------------------------------------------------------------------
def _add_header_footer(doc):
    section = doc.sections[0]

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if LOGO_FOR_DOCX.exists():
        header_p.add_run().add_picture(str(LOGO_FOR_DOCX), height=Inches(0.32))
    name = header_p.add_run("  Kanak Prabhakar")
    name.font.size = Pt(8)
    name.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead = footer_p.add_run("Converted with Likhit  ·  Page ")
    lead.font.size = Pt(8)
    lead.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    # PAGE field: Word khud har page ka number bharta hai
    run = footer_p.add_run()
    run.font.size = Pt(8)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


# ---------------------------------------------------------------------------
# Main function: pages ki list se .docx banao.
# ---------------------------------------------------------------------------
def build_document(pages, out_path, include_scans=False, page_breaks=True,
                   title=None):
    doc = Document()

    # Base font: Calibri 11 wahi hai jo Word khud naye document mein deta hai,
    # isliye file "normal" Word document jaisi lagti hai
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    # Hindi (Devanagari) text ke liye complex-script font bhi set karo,
    # warna kuch computers par boxes dikhte hain
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:cs"), "Nirmala UI")

    doc.core_properties.title = title or Path(out_path).stem
    doc.core_properties.author = "Likhit"

    _add_header_footer(doc)

    if title:
        doc.add_heading(title, level=0)

    for i, page in enumerate(pages):
        if i > 0 and page_breaks:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        # v1.1: teen haalat -> fail hua page, khaali page, ya asli text
        if page.error and not (page.text or "").strip():
            _add_note(doc, f"This page could not be read ({page.name}): {page.error}")
        elif (page.text or "").strip() == NO_TEXT:
            _add_note(doc, f"No handwriting was found on this page ({page.name}).")
        else:
            _write_text(doc, page.text or "")

        if include_scans and page.image is not None:
            cap = doc.add_paragraph()
            cap_run = cap.add_run(f"Original: {page.name}")
            cap_run.italic = True
            cap_run.font.size = Pt(8)
            cap_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            png = imaging.to_png_bytes(page.image)
            doc.add_picture(io.BytesIO(png), width=Inches(4.5))

    out_path = Path(out_path)
    if out_path.suffix.lower() != ".docx":
        out_path = out_path.with_suffix(".docx")
    doc.save(out_path)
    return out_path
